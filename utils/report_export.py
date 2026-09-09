"""Экспорт отчёта по региону в XLSX и PDF (ТЗ §13)."""

import io
import logging
import os
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.reports import RegionReport
from utils.parser import format_kopecks

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="8A6A3A")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_table(sheet, headers: list[str], rows: list[tuple], widths: list[int] | None = None) -> None:
    sheet.append(headers)
    for cell in sheet[sheet.max_row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append(list(row))
    for index, width in enumerate(widths or [], start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def build_region_xlsx(report: RegionReport) -> bytes:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Сводка"
    summary.append([f"Отчёт: {report.region_name}"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([f"Период: {report.period_label} ({report.start:%d.%m.%Y} — {report.end:%d.%m.%Y})"])
    summary.append([])
    _write_table(
        summary,
        ["Показатель", "Значение"],
        [
            ("Баланс региона (всего)", format_kopecks(report.balance)),
            ("Доходы за период", format_kopecks(report.income)),
            ("Расходы за период", format_kopecks(report.expense)),
            ("Итог за период", format_kopecks(report.income - report.expense)),
            ("Численность состава", sum(count for _, count in report.members_by_status)),
            *[(f"— {label}", count) for label, count in report.members_by_status],
            ("Мероприятий за период", len(report.events)),
        ],
        widths=[36, 22],
    )

    members_sheet = workbook.create_sheet("Состав")
    _write_table(
        members_sheet,
        ["ФИО", "Статус", "Телефон", "Дата рождения", "Вступление в Академисты", "ВУЗ", "Факультет", "Курс", "Место работы"],
        report.members,
        widths=[34, 18, 20, 16, 16, 28, 24, 10, 26],
    )

    finance_sheet = workbook.create_sheet("Финансы")
    _write_table(
        finance_sheet,
        ["Дата", "Тип", "Категория", "Мероприятие", "Сумма", "Комментарий", "Внёс"],
        [
            (tx_date.strftime("%d.%m.%Y"), tx_type, category, event, format_kopecks(amount), comment, author)
            for tx_date, tx_type, category, event, amount, comment, author in report.transactions
        ],
        widths=[12, 10, 24, 24, 14, 34, 22],
    )
    finance_sheet.append([])
    _write_table(
        finance_sheet,
        ["Расходы по категориям", "Сумма"],
        [(name, format_kopecks(amount)) for name, amount in report.expense_by_category],
    )
    finance_sheet.append([])
    _write_table(
        finance_sheet,
        ["Доходы по категориям", "Сумма"],
        [(name, format_kopecks(amount)) for name, amount in report.income_by_category],
    )

    events_sheet = workbook.create_sheet("Мероприятия")
    _write_table(
        events_sheet,
        ["Дата", "Название", "Статус", "План, ₽", "Факт, ₽", "Участников"],
        [
            (
                event_date.strftime("%d.%m.%Y"),
                title,
                status,
                format_kopecks(planned) if planned is not None else "",
                format_kopecks(fact),
                attended,
            )
            for event_date, title, status, planned, fact, attended in report.events
        ],
        widths=[12, 40, 16, 14, 14, 12],
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- PDF ---------------------------------------------------------------------

# Встроенные шрифты reportlab кириллицу не покрывают, поэтому подбираем системный
# TrueType. Путь можно задать явно через REPORT_FONT_PATH.
_FONT_CANDIDATES = [
    os.getenv("REPORT_FONT_PATH", ""),
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "C:/Windows/Fonts/segoeui.ttf",
    "/Library/Fonts/Arial.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
]
_BOLD_SUFFIXES = {"DejaVuSans.ttf": "DejaVuSans-Bold.ttf", "arial.ttf": "arialbd.ttf",
                  "LiberationSans-Regular.ttf": "LiberationSans-Bold.ttf", "segoeui.ttf": "segoeuib.ttf",
                  "Arial.ttf": "Arial Bold.ttf"}

_FONT_NAME = "Helvetica"
_FONT_BOLD = "Helvetica-Bold"
_fonts_registered = False


def _register_fonts() -> tuple[str, str]:
    global _FONT_NAME, _FONT_BOLD, _fonts_registered
    if _fonts_registered:
        return _FONT_NAME, _FONT_BOLD

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    _fonts_registered = True
    for candidate in _FONT_CANDIDATES:
        if not candidate:
            continue
        path = Path(candidate)
        if not path.is_file():
            continue
        try:
            pdfmetrics.registerFont(TTFont("ReportFont", str(path)))
            _FONT_NAME = "ReportFont"
            bold_path = path.with_name(_BOLD_SUFFIXES.get(path.name, ""))
            if bold_path.name and bold_path.is_file():
                pdfmetrics.registerFont(TTFont("ReportFont-Bold", str(bold_path)))
                _FONT_BOLD = "ReportFont-Bold"
            else:
                _FONT_BOLD = "ReportFont"
            return _FONT_NAME, _FONT_BOLD
        except Exception:  # noqa: BLE001 — битый шрифт не повод падать, пробуем следующий
            continue

    logger.warning(
        "Не найден TrueType-шрифт с кириллицей — PDF будет свёрстан Helvetica "
        "и кириллица отобразится некорректно. Укажите REPORT_FONT_PATH в .env."
    )
    return _FONT_NAME, _FONT_BOLD


def build_region_pdf(report: RegionReport) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font, font_bold = _register_fonts()

    title_style = ParagraphStyle("title", fontName=font_bold, fontSize=16, leading=20, spaceAfter=6)
    subtitle_style = ParagraphStyle("subtitle", fontName=font, fontSize=10, leading=14, textColor=colors.grey)
    heading_style = ParagraphStyle("heading", fontName=font_bold, fontSize=12, leading=16, spaceBefore=10, spaceAfter=6)

    def table(headers: list[str], rows: list[list], widths: list[float]) -> Table:
        data = [[Paragraph(f"<b>{h}</b>", ParagraphStyle("th", fontName=font_bold, fontSize=9)) for h in headers]]
        cell_style = ParagraphStyle("td", fontName=font, fontSize=9, leading=11)
        data.extend([[Paragraph(str(value), cell_style) for value in row] for row in rows])
        component = Table(data, colWidths=widths, repeatRows=1)
        component.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#8A6A3A")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D9CBB0")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F1E6")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        return component

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
        title=f"Отчёт — {report.region_name}",
    )

    total_width = doc.width
    story = [
        Paragraph(f"Отчёт регионального отделения: {report.region_name}", title_style),
        Paragraph(
            f"Период: {report.period_label} ({report.start:%d.%m.%Y} — {report.end:%d.%m.%Y})", subtitle_style
        ),
        Spacer(1, 8),
        Paragraph("Сводка", heading_style),
        table(
            ["Показатель", "Значение"],
            [
                ["Баланс региона (всего)", f"{format_kopecks(report.balance)} ₽"],
                ["Доходы за период", f"{format_kopecks(report.income)} ₽"],
                ["Расходы за период", f"{format_kopecks(report.expense)} ₽"],
                ["Итог за период", f"{format_kopecks(report.income - report.expense)} ₽"],
                *[[label, str(count)] for label, count in report.members_by_status],
                ["Мероприятий за период", str(len(report.events))],
            ],
            [total_width * 0.5, total_width * 0.5],
        ),
        Paragraph("Состав отделения", heading_style),
        table(
            ["ФИО", "Статус", "Телефон", "Дата рождения", "Вступление в Академисты", "ВУЗ", "Факультет", "Курс", "Место работы"],
            [list(row) for row in report.members] or [["—", "", "", "", "", "", "", "", ""]],
            [
                total_width * 0.18,
                total_width * 0.10,
                total_width * 0.12,
                total_width * 0.10,
                total_width * 0.10,
                total_width * 0.15,
                total_width * 0.12,
                total_width * 0.06,
                total_width * 0.07,
            ],
        ),
        PageBreak(),
        Paragraph("Финансовые операции", heading_style),
        table(
            ["Дата", "Тип", "Категория", "Мероприятие", "Сумма, ₽", "Комментарий"],
            [
                [
                    tx_date.strftime("%d.%m.%Y"),
                    tx_type,
                    category,
                    event,
                    format_kopecks(amount),
                    comment,
                ]
                for tx_date, tx_type, category, event, amount, comment, _author in report.transactions
            ]
            or [["—", "", "", "", "", ""]],
            [
                total_width * 0.1,
                total_width * 0.09,
                total_width * 0.2,
                total_width * 0.2,
                total_width * 0.12,
                total_width * 0.29,
            ],
        ),
        Paragraph("Мероприятия", heading_style),
        table(
            ["Дата", "Название", "Статус", "План, ₽", "Факт, ₽", "Участников"],
            [
                [
                    event_date.strftime("%d.%m.%Y"),
                    title,
                    status,
                    format_kopecks(planned) if planned is not None else "—",
                    format_kopecks(fact),
                    str(attended),
                ]
                for event_date, title, status, planned, fact, attended in report.events
            ]
            or [["—", "", "", "", "", ""]],
            [
                total_width * 0.11,
                total_width * 0.37,
                total_width * 0.16,
                total_width * 0.12,
                total_width * 0.12,
                total_width * 0.12,
            ],
        ),
    ]

    doc.build(story)
    return buffer.getvalue()

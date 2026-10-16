"""Regression tests for exports, expiring media, browser sessions and headcounts."""
import io
import time
from datetime import date, timedelta
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from api.auth import ensure_impersonation_read_only
from api.main import app
from api.rate_limit import _limiter
from database.models import BrowserSession, Event, NewsPhoto, NewsPost, User
from services.browser_login import COOKIE_NAME, approve, create_session, digest, pending_code
from services.recurring_events import materialize_event
from services.report_downloads import decode_ticket
from tests.conftest import login


@pytest.fixture(autouse=True)
def fresh_limits():
    _limiter._events.clear()


async def event(client, world, **fields):
    login(world["leader_moscow"])
    response = await client.post("/api/events", json={
        "region_id": world["moscow"].id, "title": "Открытая встреча",
        "date": date.today().isoformat(), "description": "Встреча клуба", **fields,
    })
    assert response.status_code == 200, response.text
    return response.json()


async def test_headcount_is_separate_and_survives_old_attendance_client(client, world):
    item = await event(client, world, total_attendees=60)
    assert item["total_attendees"] == 60
    url = f"/api/events/{item['id']}"
    assert (await client.put(url + "/attendance", json={"member_ids": []})).status_code == 200
    assert (await client.get(url)).json()["total_attendees"] == 60
    assert (await client.put(url + "/attendance", json={"member_ids": [], "total_attendees": 0})).status_code == 200
    assert (await client.get(url)).json()["total_attendees"] == 0
    await client.put(url + "/attendance", json={"member_ids": [], "total_attendees": None})
    assert (await client.get(url)).json()["total_attendees"] is None


@pytest.mark.parametrize("value", [-1, 1_000_001, 2.5, True, "60"])
async def test_invalid_headcounts_rejected(client, world, value):
    item = await event(client, world)
    response = await client.put(f"/api/events/{item['id']}/attendance", json={"member_ids": [], "total_attendees": value})
    assert response.status_code == 422


async def test_headcount_scoped_to_region(client, world):
    item = await event(client, world, total_attendees=60)
    login(world["leader_tula"])
    assert (await client.patch(f"/api/events/{item['id']}", json={"total_attendees": 99})).status_code == 403


async def test_recurrences_do_not_copy_attendance(session, world):
    template = Event(region_id=world["moscow"].id, title="Встреча", date=date.today(),
                     description="Встреча", is_recurring=True, recurrence_rule="weekly", total_attendees=60)
    session.add(template)
    await session.commit()
    children = await materialize_event(session, template)
    assert children and all(child.total_attendees is None for child in children)


async def test_export_contains_both_counts_and_downloads_without_tma(client, world, tmp_path, monkeypatch):
    monkeypatch.setattr("services.report_downloads.STORAGE_DIR", tmp_path)
    item = await event(client, world, total_attendees=60)
    member = (await client.post("/api/members", json={"region_id": world["moscow"].id, "full_name": "Участник встречи"})).json()
    await client.put(f"/api/events/{item['id']}/attendance", json={"member_ids": [member["id"]]})
    base = f"/api/reports/region/file?region_id={world['moscow'].id}&period=all"
    prepared = await client.post(base + "&format=xlsx")
    assert prepared.status_code == 200, prepared.text
    url = prepared.json()["url"]
    app.dependency_overrides.clear()
    response = await client.get(url)
    assert response.status_code == 200, response.text
    assert 'attachment; filename="report.xlsx"' in response.headers["content-disposition"]
    assert int(response.headers["content-length"]) == len(response.content)
    workbook = load_workbook(io.BytesIO(response.content))
    rows = list(workbook["Мероприятия"].values)
    assert rows[0][-2:] == ("Отмечено членов", "Всего посетителей")
    assert rows[1][-2:] == (1, 60)
    login(world["leader_moscow"])
    pdf = await client.post(base + "&format=pdf")
    assert pdf.status_code == 200, pdf.text
    assert (await client.get(pdf.json()["url"])).content.startswith(b"%PDF")


async def test_export_tickets_tampering_expiry_and_revocation(client, world, tmp_path, monkeypatch, session):
    monkeypatch.setattr("services.report_downloads.STORAGE_DIR", tmp_path)
    login(world["leader_moscow"])
    base = f"/api/reports/region/file?region_id={world['moscow'].id}"
    response = await client.post(base)
    assert response.status_code == 200
    url = response.json()["url"]
    parsed = urlsplit(url)
    token = parse_qs(parsed.query)["t"][0]
    assert (await client.get(parsed.path)).status_code == 403
    assert (await client.get(url + "x")).status_code == 403
    assert (await client.get(parsed.path.replace(".xlsx", "a.xlsx") + "?t=" + token)).status_code == 403
    login(world["leader_tula"])
    assert (await client.post(base)).status_code == 403
    login(world["cell_leader"])
    assert (await client.post(base)).status_code == 403
    login(world["leader_moscow"])
    await session.execute(update(User).where(User.id == world["leader_moscow"].id).values(is_active=False))
    await session.commit()
    assert (await client.get(url)).status_code == 403
    current = time.time()
    monkeypatch.setattr("services.report_downloads.time.time", lambda: current + 700)
    with pytest.raises(Exception) as exc:
        decode_ticket(token)
    assert exc.value.status_code == 403


async def test_unknown_headcount_is_not_reported_as_zero(client, world):
    await event(client, world)
    response = await client.get(f"/api/reports/region?region_id={world['moscow'].id}&period=all")
    rows = list(load_workbook(io.BytesIO(response.content))["Мероприятия"].values)
    assert rows[1][-2:] == (0, "Не указано")


async def test_photo_url_refresh_requires_auth_and_a_live_post(client, world, session):
    post = NewsPost(author_user_id=world["federal"].id, byline="Братство", text="Новость")
    session.add(post)
    await session.flush()
    photo = NewsPhoto(post_id=post.id, stored_path="news/test.jpg", original_name="test.jpg", content_type="image/jpeg")
    session.add(photo)
    await session.commit()
    url = f"/api/news/photos/{photo.id}/url"
    assert (await client.get(url)).status_code == 401
    login(world["leader_moscow"])
    data = (await client.get(url)).json()
    assert data["url"].startswith(f"/api/news/photos/{photo.id}?t=")
    assert (await client.get("/api/news/photos/999999/url")).status_code == 404


async def test_browser_approval_one_time_and_cookie_bound(client, world, session):
    secret, grant, code = await create_session()
    assert await pending_code(grant) == code
    assert (await client.get("/api/me")).status_code == 401
    client.cookies.set(COOKIE_NAME, secret)
    assert not (await client.get("/api/auth/browser/status")).json()["authenticated"]
    assert await approve(grant, world["leader_moscow"])
    assert not await approve(grant, world["leader_tula"])
    assert await pending_code(grant) is None
    assert (await client.get("/api/auth/browser/status")).json()["authenticated"]
    assert (await client.get("/api/me")).json()["id"] == world["leader_moscow"].id
    # Invalid signed credentials must not silently fall back to a cookie.
    assert (await client.get("/api/me", headers={"Authorization": "tma invalid"})).status_code == 401
    assert (await client.post("/api/auth/browser/logout")).status_code == 403
    assert (await client.post("/api/auth/browser/logout", headers={"Origin": "https://evil.invalid"})).status_code == 403
    assert (await client.post("/api/auth/browser/logout", headers={"Origin": "http://test"})).status_code == 200
    assert (await client.get("/api/me")).status_code == 401


async def test_browser_start_cookie_options_and_origin(client, world, monkeypatch):
    class FakeBot:
        async def get_me(self): return SimpleNamespace(username="test_bot")
    monkeypatch.setattr("api.routers.browser_login.get_notifier_bot", lambda: FakeBot())
    assert (await client.post("/api/auth/browser/start")).status_code == 403
    response = await client.post("/api/auth/browser/start", headers={"Origin": "http://test"})
    assert response.status_code == 200, response.text
    assert "HttpOnly" in response.headers["set-cookie"]
    assert "SameSite=strict" in response.headers["set-cookie"]
    assert response.json()["url"].startswith("https://t.me/test_bot?start=web_")
    assert len(response.json()["code"]) == 6


async def test_expired_and_inactive_browser_sessions_are_denied(client, world, session):
    secret, grant, _ = await create_session()
    await session.execute(update(BrowserSession).where(BrowserSession.secret_hash == digest(secret)).values(expires_at=1))
    await session.commit()
    assert not await approve(grant, world["leader_moscow"])
    client.cookies.set(COOKIE_NAME, secret)
    assert (await client.get("/api/me")).status_code == 401
    secret, grant, _ = await create_session()
    assert await approve(grant, world["leader_moscow"])
    client.cookies.set(COOKIE_NAME, secret)
    await session.execute(update(User).where(User.id == world["leader_moscow"].id).values(is_active=False))
    await session.commit()
    assert (await client.get("/api/me")).status_code == 401


async def test_browser_cookie_mutations_require_origin(client, world):
    secret, grant, _ = await create_session()
    assert await approve(grant, world["leader_moscow"])
    client.cookies.set(COOKIE_NAME, secret)
    data = {"region_id": world["moscow"].id, "full_name": "Новый Участник"}
    assert (await client.post("/api/members", json=data)).status_code == 403
    assert (await client.post("/api/members", json=data, headers={"Origin": "https://evil.invalid"})).status_code == 403
    assert (await client.post("/api/members", json=data, headers={"Origin": "http://test"})).status_code == 200


async def test_browser_idempotency_is_scoped_to_the_session(client, world):
    url = "/api/members"
    headers = {"Origin": "http://test", "Idempotency-Key": "browser-session-test-same-key"}
    for role, region, name in [("leader_moscow", "moscow", "Московский Участник"),
                               ("leader_tula", "tula", "Тульский Участник")]:
        secret, grant, _ = await create_session()
        assert await approve(grant, world[role])
        client.cookies.set(COOKIE_NAME, secret)
        response = await client.post(url, json={"region_id": world[region].id, "full_name": name}, headers=headers)
        assert response.status_code == 200, response.text
        assert response.json()["full_name"] == name


async def test_excluding_a_member_revokes_browser_sessions(client, world, session):
    from database.models import Member
    from services.admin_actions import exclude_member
    member = Member(region_id=world["moscow"].id, full_name="Удаляемый Участник")
    session.add(member)
    await session.flush()
    user = User(telegram_id=77777, full_name=member.full_name, role="participant", member_id=member.id)
    session.add(user)
    await session.commit()
    secret, grant, _ = await create_session()
    assert await approve(grant, user)
    client.cookies.set(COOKIE_NAME, secret)
    assert (await client.get("/api/me")).status_code == 200
    await exclude_member(session, member.id)
    assert (await client.get("/api/me")).status_code == 401
    assert not (await session.execute(select(BrowserSession).where(BrowserSession.secret_hash == digest(secret)))).first()


def test_read_only_impersonation_can_prepare_export_but_not_mutate(world):
    user = world["leader_moscow"]
    user._impersonated_by = 1
    ensure_impersonation_read_only(user, "POST", "/api/reports/region/file")
    with pytest.raises(Exception) as error:
        ensure_impersonation_read_only(user, "POST", "/api/events")
    assert error.value.status_code == 403


async def test_migration_preserves_old_events_and_can_run_twice(tmp_path, monkeypatch):
    from sqlalchemy import text
    from sqlalchemy.ext.asyncio import create_async_engine
    from scripts import migrate_service_reliability as migration
    engine = create_async_engine(f"sqlite+aiosqlite:///{(tmp_path / 'old.db').as_posix()}")
    async with engine.begin() as connection:
        await connection.execute(text("CREATE TABLE events (id INTEGER PRIMARY KEY, title TEXT)"))
        await connection.execute(text("INSERT INTO events VALUES (1, 'Old event')"))
        await connection.execute(text("CREATE TABLE users (id INTEGER PRIMARY KEY)"))
    monkeypatch.setattr(migration, "engine", engine)
    await migration.migrate()
    await migration.migrate()
    async with engine.connect() as connection:
        assert (await connection.execute(text("SELECT title, total_attendees FROM events"))).one() == ("Old event", None)
        assert (await connection.execute(text("SELECT count(*) FROM browser_sessions"))).scalar() == 0
    await engine.dispose()


async def test_finance_download_is_scoped_and_supports_head(client, world, tmp_path, monkeypatch):
    monkeypatch.setattr("services.report_downloads.STORAGE_DIR", tmp_path)
    login(world["leader_moscow"])
    response = await client.post("/api/downloads/prepare", json={"kind": "finance", "region_id": world["moscow"].id})
    assert response.status_code == 200, response.text
    url = response.json()["url"]
    app.dependency_overrides.clear()
    response = await client.get(url)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    head = await client.head(url)
    assert head.status_code == 200 and head.content == b""
    assert head.headers["content-length"] == response.headers["content-length"]
    login(world["leader_tula"])
    assert (await client.post("/api/downloads/prepare", json={"kind": "finance", "region_id": world["moscow"].id})).status_code == 403


async def test_document_download_revoked_when_document_is_deleted(client, world, tmp_path, monkeypatch):
    monkeypatch.setattr("services.report_downloads.STORAGE_DIR", tmp_path)
    login(world["leader_moscow"])
    upload = await client.post("/api/documents", data={"region_id": world["moscow"].id, "title": "Памятка"},
                               files={"file": ("Памятка.txt", "Проверка документа".encode(), "text/plain")})
    assert upload.status_code == 200, upload.text
    document_id = upload.json()["id"]
    response = await client.post("/api/downloads/prepare", json={"kind": "document", "document_id": document_id})
    assert response.status_code == 200, response.text
    url = response.json()["url"]
    data = await client.get(url)
    assert data.status_code == 200, data.text
    assert data.content.decode() == "Проверка документа"
    login(world["leader_tula"])
    assert (await client.post("/api/downloads/prepare", json={"kind": "document", "document_id": document_id})).status_code == 403
    login(world["leader_moscow"])
    assert (await client.delete(f"/api/documents/{document_id}")).status_code == 200
    assert (await client.get(url)).status_code == 404


async def test_export_limit_covers_legacy_and_new_routes(client, world):
    from services.report_downloads import generation_slots
    login(world["leader_moscow"])
    await generation_slots.acquire()
    await generation_slots.acquire()
    try:
        region_id = world["moscow"].id
        assert (await client.get(f"/api/reports/region?region_id={region_id}")).status_code == 429
        assert (await client.post(f"/api/reports/region/file?region_id={region_id}")).status_code == 429
        assert (await client.post("/api/downloads/prepare", json={"kind": "finance", "region_id": region_id})).status_code == 429
        assert (await client.get("/api/me")).status_code == 200
    finally:
        generation_slots.release()
        generation_slots.release()
